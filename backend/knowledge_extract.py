import io,base64
from pathlib import PurePosixPath
from fastapi import HTTPException

TEXT_TYPES={'md','markdown','txt','csv','json','yaml','yml'}
OFFICE_TYPES={'docx','xlsx','pptx','pdf'}
SUPPORTED=TEXT_TYPES|OFFICE_TYPES
MAX_BYTES=25*1024*1024
MAX_TEXT=400000

def extension(path):return PurePosixPath(path).suffix.lower().lstrip('.')
def decode_payload(content,encoding):
    if encoding=='text':return content.encode()
    try:return base64.b64decode(content,validate=True)
    except Exception:raise HTTPException(422,'File content must be valid base64')
def extract_text(path,raw):
    ext=extension(path)
    if ext not in SUPPORTED:raise HTTPException(415,f'Unsupported file type .{ext}. Supported: Markdown, text, CSV, JSON, YAML, Word (.docx), Excel (.xlsx), PowerPoint (.pptx), PDF')
    if len(raw)>MAX_BYTES:raise HTTPException(413,'Files must be 25 MB or smaller')
    try:
        if ext in TEXT_TYPES:text=raw.decode('utf-8-sig',errors='replace')
        elif ext=='docx':text=_docx(raw)
        elif ext=='xlsx':text=_xlsx(raw)
        elif ext=='pptx':text=_pptx(raw)
        else:text=_pdf(raw)
    except HTTPException:raise
    except Exception:raise HTTPException(422,f'Could not read {PurePosixPath(path).name}. The file may be corrupted or password protected')
    text='\n'.join(line.rstrip() for line in text.splitlines()).strip()
    if not text:raise HTTPException(422,f'{PurePosixPath(path).name} contains no readable text')
    return text[:MAX_TEXT]
def _docx(raw):
    import docx
    document=docx.Document(io.BytesIO(raw));parts=[]
    for block in document.paragraphs:
        style=(block.style.name or '').lower()
        prefix='# ' if style.startswith('heading 1') or style=='title' else '## ' if style.startswith('heading 2') else '### ' if style.startswith('heading') else ''
        if block.text.strip():parts.append(prefix+block.text.strip())
    for table in document.tables:
        for row in table.rows:parts.append(' | '.join(cell.text.strip() for cell in row.cells))
    return '\n'.join(parts)
def _xlsx(raw):
    import openpyxl
    workbook=openpyxl.load_workbook(io.BytesIO(raw),read_only=True,data_only=True);parts=[]
    for sheet in workbook.worksheets:
        parts.append(f'## Sheet: {sheet.title}')
        for row in sheet.iter_rows(values_only=True):
            cells=['' if value is None else str(value) for value in row]
            if any(cells):parts.append(' | '.join(cells))
    return '\n'.join(parts)
def _pptx(raw):
    from pptx import Presentation
    deck=Presentation(io.BytesIO(raw));parts=[]
    for index,slide in enumerate(deck.slides,1):
        parts.append(f'## Slide {index}')
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():parts.append(shape.text_frame.text.strip())
    return '\n'.join(parts)
def _pdf(raw):
    from pypdf import PdfReader
    reader=PdfReader(io.BytesIO(raw))
    if reader.is_encrypted:raise HTTPException(422,'Encrypted PDF files are not supported')
    return '\n'.join((page.extract_text() or '') for page in reader.pages[:300])
